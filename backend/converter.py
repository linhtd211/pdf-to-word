from __future__ import annotations

import io
import logging
import re
import unicodedata
from pathlib import Path
from typing import Any, Callable, cast

import fitz
import numpy as np
import pdfplumber
import pytesseract
from PIL import Image, ImageChops, ImageOps
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
try:
    from pdf2docx import Converter
except ImportError:
    Converter = None  # Optional fast-path; native editable conversion remains available.
from .page_policy import PagePolicy, apply_page_setup, resolve_page_setups
from .paragraphs import can_join_text_blocks, merge_adjacent_text_events, split_lines_into_paragraphs
from .quality import inspect_conversion

logger = logging.getLogger(__name__)
ProgressCallback = Callable[[int, str], None]


class ConversionError(Exception):
    """Lỗi có thể hiển thị an toàn cho người dùng."""


def inspect_pdf(path: Path, max_pages: int = 100) -> dict:
    try:
        pdf: Any = fitz.open(path)
    except Exception as exc:
        raise ConversionError("Không thể đọc PDF. Tệp có thể bị hỏng hoặc không đúng định dạng.") from exc
    try:
        if pdf.needs_pass:
            raise ConversionError("PDF đang được bảo vệ bằng mật khẩu. Vui lòng mở khóa rồi tải lại.")
        count = len(pdf)
        if count < 1:
            raise ConversionError("PDF không có trang để chuyển đổi.")
        if count > max_pages:
            raise ConversionError(f"PDF có {count} trang; giới hạn hiện tại là {max_pages} trang.")
        page_states = []
        for page_index in range(count):
            page = pdf[page_index]
            text = page.get_text("text").strip()
            has_images = bool(page.get_images(full=True))
            page_states.append("text" if text else "scan" if has_images else "blank")
        text_pages = page_states.count("text")
        scan_pages = page_states.count("scan")
        return {
            "pages": count,
            "text_pages": text_pages,
            "scan_pages": scan_pages,
            "kind": "electronic" if scan_pages == 0 else "scanned" if text_pages == 0 else "mixed",
            "page_sizes": [(float(pdf[i].rect.width), float(pdf[i].rect.height)) for i in range(count)],
        }
    finally:
        pdf.close()


def _progress(callback: ProgressCallback | None, value: int, label: str) -> None:
    if callback:
        callback(value, label)


def _extract_page_number_line(blocks: list[dict], page_index: int, page_width: float, page_height: float) -> bool:
    """Remove an isolated centered page number from body flow so it cannot split a sentence."""
    found = False
    for block in blocks:
        if block.get("type") != 0:
            continue
        lines = block.get("lines", [])
        kept = []
        for line in lines:
            text = _line_text(line)
            bbox = line.get("bbox", (0, 0, 0, 0))
            center = (float(bbox[0]) + float(bbox[2])) / 2
            is_page_number = (
                text == str(page_index + 1)
                and abs(center - page_width / 2) <= page_width * 0.07
                and float(bbox[1]) <= min(45.0, page_height * 0.06)
            )
            if is_page_number:
                found = True
            else:
                kept.append(line)
        if len(kept) != len(lines):
            block["lines"] = kept
            boxes = [line["bbox"] for line in kept if line.get("bbox")]
            if boxes:
                block["bbox"] = (
                    min(float(box[0]) for box in boxes),
                    min(float(box[1]) for box in boxes),
                    max(float(box[2]) for box in boxes),
                    max(float(box[3]) for box in boxes),
                )
    return found


def _add_page_number_field(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    if paragraph._p.xpath(".//w:instrText"):
        return
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _join_word_paragraphs(left, right) -> None:
    """Append runs from a continuation paragraph and remove its redundant paragraph boundary."""
    left.add_run(" ")
    for child in list(right._p):
        if child.tag != qn("w:pPr"):
            left._p.append(child)
    parent = right._p.getparent()
    if parent is not None:
        parent.remove(right._p)


def _plain_upper(text: str) -> str:
    decomposed = unicodedata.normalize("NFD", text).replace("Đ", "D").replace("đ", "d")
    return "".join(char for char in decomposed if unicodedata.category(char) != "Mn").upper()


def _line_text(line: dict) -> str:
    return "".join(str(span.get("text", "")) for span in line.get("spans", [])).strip()


def _detect_administrative_header(source: Path) -> dict | None:
    """Recognize the first-page national header and its parallel left/right blocks."""
    try:
        with fitz.open(source) as pdf:
            if not len(pdf):
                return None
            page: Any = pdf[0]
            page_width = float(page.rect.width)
            lines: list[dict] = []
            for block in page.get_text("dict").get("blocks", []):
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    text = _line_text(line)
                    if text:
                        lines.append({
                            "text": text,
                            "norm": _plain_upper(text),
                            "bbox": tuple(float(x) for x in line["bbox"]),
                            "spans": line.get("spans", []),
                        })
    except Exception:
        logger.exception("Could not inspect the first-page administrative header")
        return None

    national = next((line for line in lines if "CONG HOA XA HOI CHU NGHIA VIET NAM" in line["norm"]), None)
    tagline = next((line for line in lines if "DOC LAP" in line["norm"] and "TU DO" in line["norm"] and "HANH PHUC" in line["norm"]), None)
    if national is None or tagline is None:
        return None

    midpoint = page_width * 0.5
    header_lines = [national, tagline]
    consumed = {national["norm"], tagline["norm"]}
    agency: list[dict] = []
    number: list[dict] = []
    subject: list[dict] = []
    date: list[dict] = []
    top_limit = max(national["bbox"][3], tagline["bbox"][3]) + 44

    # PDF forms may place the filled number/day/month as independent text objects
    # over blank placeholders on the number/date row.
    numeric_lines = [line for line in lines if re.fullmatch(r"\d{1,6}", line["text"].strip())]
    numeric_groups: list[list[dict]] = []
    for line in sorted(numeric_lines, key=lambda item: (item["bbox"][1], item["bbox"][0])):
        if numeric_groups and abs(line["bbox"][1] - numeric_groups[-1][0]["bbox"][1]) <= 2.5:
            numeric_groups[-1].append(line)
        else:
            numeric_groups.append([line])
    number_value = None
    date_values: list[dict] = []
    filled_value_lines: list[dict] = []
    for group in numeric_groups:
        left_values = [line for line in group if (line["bbox"][0] + line["bbox"][2]) / 2 < midpoint]
        right_values = sorted(
            [line for line in group if (line["bbox"][0] + line["bbox"][2]) / 2 >= midpoint],
            key=lambda item: item["bbox"][0],
        )
        if len(left_values) == 1 and len(right_values) >= 2 and group[0]["bbox"][1] <= top_limit:
            number_value = left_values[0]["text"].strip()
            date_values = right_values[:2]
            filled_value_lines = [*left_values, *date_values]
            break

    for line in lines:
        if line is national or line is tagline or line in filled_value_lines:
            continue
        x0, y0, x1, _y1 = line["bbox"]
        if y0 > top_limit:
            continue
        normalized = line["norm"].lstrip()
        is_date = ("NGAY" in normalized and ("THANG" in normalized or "NAM" in normalized))
        if is_date and (x0 + x1) / 2 >= midpoint - 8:
            date.append(line)
        elif normalized.startswith(("SO:", "NO:")) and (x0 + x1) / 2 < midpoint:
            number.append(line)
        elif normalized.startswith(("V/V", "VE VIEC", "TRICH YEU")) and (x0 + x1) / 2 < midpoint:
            subject.append(line)
        elif (x0 + x1) / 2 < midpoint and y0 <= national["bbox"][1] + 30:
            agency.append(line)

    # A wrapped V/v line is still part of the header subject, not the body.
    if subject:
        subject.sort(key=lambda line: (line["bbox"][1], line["bbox"][0]))
        previous_y = subject[-1]["bbox"][1]
        while True:
            candidates = [line for line in lines if line not in header_lines and line not in filled_value_lines and line not in subject and line["bbox"][1] > previous_y]
            if not candidates:
                break
            candidate = min(candidates, key=lambda line: (line["bbox"][1], line["bbox"][0]))
            x0, y0, x1, _y1 = candidate["bbox"]
            if y0 - previous_y > 20 or (x0 + x1) / 2 >= midpoint:
                break
            subject.append(candidate)
            previous_y = y0

    if number_value and number:
        for line in number:
            if "SO:" in line["norm"] and not re.search(r"\b\d{3,6}\b", line["text"]):
                original_norm = line["norm"]
                line["text"] = re.sub(
                    r"(?i)(Số\s*:\s*)(?:\d{3,6}\s*)?(?=/)",
                    lambda match: f"{match.group(1).rstrip()} {number_value} ",
                    line["text"],
                    count=1,
                )
                line["norm"] = _plain_upper(line["text"])
                consumed.add(original_norm)
    if len(date_values) == 2 and date:
        day, month = (value["text"].strip() for value in date_values)
        for line in date:
            original_norm = line["norm"]
            line["text"] = re.sub(
                r"(?i)(ngày)\s*(\d{1,2})?\s*(tháng)\s*(\d{1,2})?\s*(năm)",
                lambda match: f"{match.group(1)} {match.group(2) or day} {match.group(3)} {match.group(4) or month} {match.group(5)}",
                line["text"],
                count=1,
            )
            line["norm"] = _plain_upper(line["text"])
            consumed.add(original_norm)

    for group in (agency, number, subject, date):
        consumed.update(line["norm"] for line in group)
        header_lines.extend(group)
    header_lines.extend(filled_value_lines)
    consumed.update(line["norm"] for line in filled_value_lines)

    return {
        "agency": sorted(agency, key=lambda line: (line["bbox"][1], line["bbox"][0])),
        "number": sorted(number, key=lambda line: (line["bbox"][1], line["bbox"][0])),
        "subject": sorted(subject, key=lambda line: (line["bbox"][1], line["bbox"][0])),
        "national": national,
        "tagline": tagline,
        "date": sorted(date, key=lambda line: (line["bbox"][1], line["bbox"][0])),
        "consumed": consumed,
        "bottom": max(line["bbox"][3] for line in header_lines),
    }


def convert_pdf(
    source: Path,
    output: Path,
    mode: str,
    progress: ProgressCallback | None = None,
    page_policy: PagePolicy = "source",
) -> dict:
    if mode not in {"layout", "editable"}:
        raise ConversionError("Chế độ chuyển đổi không hợp lệ.")
    if page_policy not in {"source", "administrative"}:
        raise ConversionError("Chế độ khổ giấy và thể thức không hợp lệ.")

    info = inspect_pdf(source)
    _progress(progress, 15, "Đọc cấu trúc PDF")
    _progress(progress, 30, "Phân tích bố cục")
    _progress(progress, 45, "Nhận diện văn bản và bảng biểu")

    administrative_header = _detect_administrative_header(source)
    if (
        info["kind"] == "electronic"
        and mode == "layout"
        and page_policy == "source"
        and administrative_header is None
        and Converter is not None
    ):
        try:
            converter = Converter(str(source))
            try:
                converter.convert(str(output), start=0)
            finally:
                converter.close()
        except Exception as exc:
            logger.exception("pdf2docx layout conversion failed")
            raise ConversionError("Không thể tái tạo bố cục PDF này. Hãy thử chế độ Ưu tiên chỉnh sửa.") from exc
        _progress(progress, 85, "Tạo tài liệu Word")
        _verify_docx(output)
        _progress(progress, 92, "Kiểm tra chất lượng Word")
        quality = inspect_conversion(source, output, page_policy=page_policy, mode=mode)
        return {"pages": info["pages"], "kind": info["kind"], "warnings": [], "quality": quality}

    warnings: list[str] = []
    ocr_page_metrics: list[dict[str, Any]] = []
    if administrative_header:
        warnings.append("Đã chuẩn hóa phần đầu văn bản theo Phụ lục I Nghị định 30/2020/NĐ-CP bằng bảng Word 2 cột; hãy kiểm tra lại thông tin cơ quan, số/ký hiệu và ngày ban hành.")
    is_administrative = page_policy == "administrative" or administrative_header is not None
    if page_policy == "administrative" and not administrative_header:
        warnings.append("Đã áp dụng khổ giấy, lề và font theo chế độ chuẩn hành chính; chưa nhận diện được đầy đủ Quốc hiệu–Tiêu ngữ để tự dựng header, hãy rà soát thể thức trước khi ban hành.")
    if info["scan_pages"]:
        warnings.append("PDF có trang scan: văn bản được nhận dạng bằng OCR; hãy rà soát lại dấu và bố cục phức tạp.")
    if mode == "layout" and info["kind"] != "electronic":
        warnings.append("PDF có trang scan hoặc kết hợp: chế độ Ưu tiên bố cục chưa thể tái tạo đầy đủ vị trí của trang scan; các trang này được xuất thành văn bản Word nhận dạng bằng OCR.")

    document = Document()
    normal = cast(Any, document.styles["Normal"])
    normal.font.name = "Times New Roman" if is_administrative else "Arial"
    normal.font.size = Pt(13 if is_administrative else 10.5)
    normal.paragraph_format.space_after = Pt(6 if is_administrative else 3)
    normal.paragraph_format.line_spacing = 1.0

    with fitz.open(source) as pdf, pdfplumber.open(source) as plumber_pdf:
        page_count = len(pdf)
        page_setups = resolve_page_setups(pdf, page_policy, administrative_header is not None)
        previous_setup = None
        previous_page_tail: tuple[dict, Any, Any] | None = None
        previous_table_tail: dict[str, Any] | None = None
        for page_index in range(page_count):
            page: Any = cast(Any, pdf[page_index])
            page_text = page.get_text("text").strip()
            page_images = bool(page.get_images(full=True))
            if not page_text and not page_images:
                warnings.append(f"Trang PDF {page_index + 1} trống, đã bỏ qua để tránh tạo trang Word trắng.")
                previous_page_tail = None
                previous_table_tail = None
                continue

            setup = page_setups[page_index]
            if previous_setup is None:
                section = document.sections[-1]
                apply_page_setup(section, setup)
            elif setup.key != previous_setup.key:
                document.add_section(WD_SECTION.NEW_PAGE)
                section = document.sections[-1]
                apply_page_setup(section, setup)
            else:
                section = document.sections[-1]
                if mode == "layout" and page_policy == "source":
                    document.add_page_break()
            if page_index == 0 and administrative_header:
                _append_administrative_header(document, section, administrative_header)
            previous_setup = setup

            if not page_text and page_images:
                ocr_metrics = _append_ocr_page(document, page, page_index, page_count, progress)
                ocr_page_metrics.append({"page": page_index + 1, **ocr_metrics})
                confidence = ocr_metrics.get("mean_confidence")
                if confidence is not None and confidence < 55:
                    warnings.append(
                        f"Độ tin cậy OCR trang {page_index + 1} thấp ({confidence:.0f}/100); hãy đọc soát dấu tiếng Việt và nội dung trong bảng."
                    )
                previous_page_tail = None
                previous_table_tail = None
                continue

            try:
                plumber_page = plumber_pdf.pages[page_index]
                table_events = _extract_tables(plumber_page)
            except Exception:
                logger.exception("Could not inspect tables on page %s", page_index + 1)
                table_events = []
            blocks = page.get_text("dict", sort=True).get("blocks", [])
            if mode == "editable" and _extract_page_number_line(
                blocks, page_index, page.rect.width, page.rect.height
            ):
                _add_page_number_field(section)
            for table in table_events:
                table["font_size"] = _estimate_table_font_size(blocks, table["bbox"])
                table["cell_styles"] = _extract_table_cell_styles(page, blocks, table.get("cell_bboxes", []))
                first_row_styles = table["cell_styles"][0] if table["cell_styles"] else []
                if first_row_styles and sum(bool(style.get("bold")) for style in first_row_styles) >= len(first_row_styles) / 2:
                    table["header_row"] = True
            events: list[tuple[float, int, dict]] = []
            footer = _detect_administrative_footer(blocks, page.rect.width, page.rect.height) if is_administrative else None
            footer_block_ids = footer["block_ids"] if footer else set()
            for block in blocks:
                if id(block) in footer_block_ids:
                    continue
                if page_index == 0 and administrative_header and block.get("type") == 0:
                    remaining_lines = []
                    for line in block.get("lines", []):
                        line_bbox = line.get("bbox", (0, 0, 0, 0))
                        if float(line_bbox[1]) <= administrative_header["bottom"] + 2 and _plain_upper(_line_text(line)) in administrative_header["consumed"]:
                            continue
                        remaining_lines.append(line)
                    if not remaining_lines:
                        continue
                    if len(remaining_lines) != len(block.get("lines", [])):
                        block = {**block, "lines": remaining_lines}
                        line_boxes = [line["bbox"] for line in remaining_lines]
                        block["bbox"] = (
                            min(float(box[0]) for box in line_boxes),
                            min(float(box[1]) for box in line_boxes),
                            max(float(box[2]) for box in line_boxes),
                            max(float(box[3]) for box in line_boxes),
                        )
                if block.get("type") == 0:
                    block = _filter_text_block_outside_tables(block, table_events)
                    if block is None:
                        continue
                bbox = block.get("bbox", (0, 0, 0, 0))
                if block.get("type") == 1 and _inside_any_table(bbox, table_events):
                    continue
                if block.get("type") == 1 and block.get("image"):
                    events.append((float(bbox[1]), 1, block))
                elif block.get("type") == 0 and any(
                    str(span.get("text", "")).strip()
                    for line in block.get("lines", [])
                    for span in line.get("spans", [])
                ):
                    events.append((float(bbox[1]), 1, block))
            for table in table_events:
                events.append((float(table["bbox"][1]), 0, table))
            if footer:
                events.append((float(footer["bbox"][1]), 2, footer))
            previous_bottom = (
                float(administrative_header["bottom"])
                if page_index == 0 and administrative_header
                else setup.top_pt
            )
            image_order = 0
            ordered_events = sorted(events, key=lambda x: (x[0], x[1]))
            if mode == "editable":
                ordered_events = merge_adjacent_text_events(ordered_events, page.rect.width)
            first_event = ordered_events[0] if ordered_events else None
            last_event = ordered_events[-1] if ordered_events else None
            first_event_text_block = (
                first_event[2] if first_event and first_event[1] == 1 and first_event[2].get("type") == 0 else None
            )
            last_event_text_block = (
                last_event[2] if last_event and last_event[1] == 1 and last_event[2].get("type") == 0 else None
            )
            first_event_paragraph = None
            last_event_paragraph = None
            page_table_tail: dict[str, Any] | None = None
            for top, kind, item in ordered_events:
                if kind != 0 and item.get("type") == 1:
                    _append_image(document, item, page.rect.width, page.rect.height, image_order)
                    image_order += 1
                    continue
                if mode == "layout" and page_policy == "source" and kind != 1:
                    # Preserve separation around non-text objects, not every PDF text coordinate.
                    _append_vertical_gap(document, min(18.0, max(0.0, top - previous_bottom)))
                if kind == 0:
                    section_width = cast(Any, section.page_width).inches
                    section_left = cast(Any, section.left_margin).inches
                    section_right = cast(Any, section.right_margin).inches
                    available_width = (section_width - section_left - section_right) * 72
                    rows = item["rows"]
                    merge_ranges = item.get("merge_ranges", [])
                    existing_table = None
                    if (mode == "editable" or page_policy == "administrative") and previous_table_tail is not None and _can_continue_table(
                        previous_table_tail, item, page_index, float(page.rect.width), float(page.rect.height)
                    ):
                        existing_table = previous_table_tail["word_table"]
                        previous_source_table = previous_table_tail["source_table"]
                        if _same_table_header(previous_source_table, item):
                            rows = rows[1:]
                            merge_ranges = [(row_index - 1, start, end) for row_index, start, end in merge_ranges if row_index > 0]
                    word_table = _append_table(
                        document,
                        rows,
                        administrative=is_administrative,
                        column_widths=item.get("column_widths"),
                        available_width_pt=available_width,
                        source_font_size=item.get("font_size"),
                        preserve_pdf_line_breaks=False,
                        cell_styles=item.get("cell_styles"),
                        merge_ranges=merge_ranges,
                        existing_table=existing_table,
                        repeat_header=bool(item.get("header_row") and existing_table is None),
                    )
                    page_table_tail = {
                        "word_table": word_table,
                        "source_table": item,
                        "page_index": page_index,
                        "page_width": float(page.rect.width),
                        "page_height": float(page.rect.height),
                    }
                elif kind == 2:
                    _append_administrative_footer(document, item, section)
                else:
                    section_left = cast(Any, section.left_margin).inches * 72
                    section_right = cast(Any, section.right_margin).inches * 72
                    paragraphs = _append_text_block(
                        document,
                        item,
                        page.rect.width,
                        administrative=is_administrative,
                        section_left_pt=section_left,
                        section_right_pt=section_right,
                        preserve_pdf_line_breaks=False,
                    )
                    if item is first_event_text_block and paragraphs:
                        first_event_paragraph = paragraphs[0]
                    if item is last_event_text_block and paragraphs:
                        last_event_paragraph = paragraphs[-1]
                bbox = item.get("bbox", (0, 0, 0, top))
                previous_bottom = max(previous_bottom, float(bbox[3]))

            if (
                mode == "editable"
                and previous_page_tail is not None
                and first_event_text_block is not None
                and first_event_paragraph is not None
                and previous_page_tail[2].key == setup.key
                and can_join_text_blocks(
                    previous_page_tail[0],
                    first_event_text_block,
                    page_width=page.rect.width,
                    page_height=setup.height_pt,
                    across_pages=True,
                )
            ):
                _join_word_paragraphs(previous_page_tail[1], first_event_paragraph)

            if mode == "editable" and last_event_text_block is not None and last_event_paragraph is not None:
                previous_page_tail = (last_event_text_block, last_event_paragraph, setup)
            else:
                previous_page_tail = None
            table_indices = [index for index, event in enumerate(ordered_events) if event[1] == 0]
            if (mode == "editable" or page_policy == "administrative") and page_table_tail is not None and table_indices:
                last_table_index = table_indices[-1]
                if not ordered_events[last_table_index + 1 :] or _table_has_only_side_annotations_after(
                    ordered_events, last_table_index, page_table_tail["source_table"]
                ):
                    previous_table_tail = page_table_tail
                else:
                    previous_table_tail = None
            else:
                previous_table_tail = None
            if page_index < page_count - 1:
                _progress(progress, 50 + int(25 * (page_index + 1) / page_count), "Nhận diện văn bản và bảng biểu")

    _progress(progress, 85, "Tạo tài liệu Word")
    try:
        document.save(str(output))
    except Exception as exc:
        raise ConversionError("Không thể tạo tệp Word từ tài liệu này.") from exc
    _verify_docx(output)
    _progress(progress, 92, "Kiểm tra chất lượng Word")
    quality = inspect_conversion(
        source, output, page_policy=page_policy, mode=mode, ocr_pages=ocr_page_metrics
    )
    return {"pages": info["pages"], "kind": info["kind"], "warnings": warnings, "quality": quality}


def _set_administrative_page_geometry(section, width_pt: float, height_pt: float, page=None) -> None:
    """Preserve source dimensions; apply Nghị định 30 margins to standard A4 portrait pages."""
    section.page_width = Inches(max(3.0, width_pt / 72))
    section.page_height = Inches(max(3.0, height_pt / 72))
    section.orientation = WD_ORIENT.LANDSCAPE if width_pt > height_pt else WD_ORIENT.PORTRAIT

    is_a4_portrait = (
        width_pt <= height_pt
        and abs(width_pt - 595.2756) <= 2
        and abs(height_pt - 841.8898) <= 2
    )
    if is_a4_portrait:
        section.top_margin = Inches(2.0 / 2.54)
        section.bottom_margin = Inches(2.5 / 2.54)
        section.left_margin = Inches(3.0 / 2.54)
        section.right_margin = Inches(1.5 / 2.54)
        return

    if page is not None:
        boxes = [
            tuple(float(value) for value in block["bbox"])
            for block in page.get_text("dict").get("blocks", [])
            if block.get("type") == 0 and block.get("bbox")
        ]
        try:
            boxes.extend(tuple(float(value) for value in table.bbox) for table in page.find_tables().tables)
        except Exception:
            logger.debug("Could not derive landscape table margins from source geometry", exc_info=True)
        if boxes:
            left_edge = min(box[0] for box in boxes)
            top_edge = min(box[1] for box in boxes)
            right_edge = max(box[2] for box in boxes)
            bottom_edge = max(box[3] for box in boxes)
            section.left_margin = Inches(max(0.15, left_edge / 72))
            section.right_margin = Inches(max(0.15, (width_pt - right_edge) / 72))
            section.top_margin = Inches(max(0.15, top_edge / 72))
            section.bottom_margin = Inches(max(0.15, (height_pt - bottom_edge) / 72))
            return

    section.top_margin = Inches(2.0 / 2.54)
    section.bottom_margin = Inches(2.5 / 2.54)
    section.left_margin = Inches(3.0 / 2.54)
    section.right_margin = Inches(1.5 / 2.54)


def _set_paragraph_bottom_rule(paragraph, space: int = 3) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    p_pr.append(borders)


def _set_cell_margins(cell, top: int = 40, start: int = 60, bottom: int = 40, end: int = 60) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders_none(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def _append_rule_paragraph(cell, width_cm: float) -> None:
    """Add a centered, narrow rule without reducing the width of the preceding text."""
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    if cell.width is not None:
        available_cm = cell.width.inches * 2.54
        inset_cm = max(0.0, (available_cm - width_cm) / 2)
        paragraph.paragraph_format.left_indent = Inches(inset_cm / 2.54)
        paragraph.paragraph_format.right_indent = Inches(inset_cm / 2.54)
    run = paragraph.add_run("\u00a0")
    run.font.name = "Times New Roman"
    run.font.size = Pt(1)
    run.font.color.rgb = RGBColor(0, 0, 0)
    _set_paragraph_bottom_rule(paragraph, space=0)


def _append_cell_line(
    cell,
    text: str,
    size: float,
    *,
    bold: bool = False,
    italic: bool = False,
    bottom_rule: bool = False,
    rule_width_cm: float | None = None,
    character_scale: int | None = None,
    space_before: float = 0,
) -> None:
    paragraph = cell.paragraphs[0] if not cell.paragraphs[0].text and len(cell.paragraphs) == 1 else cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    if run._element.rPr is not None:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if character_scale is not None:
        scale = OxmlElement("w:w")
        scale.set(qn("w:val"), str(character_scale))
        run._element.get_or_add_rPr().append(scale)
    if bottom_rule:
        if rule_width_cm is None:
            _set_paragraph_bottom_rule(paragraph)
        else:
            _append_rule_paragraph(cell, rule_width_cm)


def _append_administrative_header(document: DocxDocument, section, header: dict) -> None:
    """Build the legal parallel header as one balanced, borderless two-column Word table."""
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    available_width = section.page_width - section.left_margin - section.right_margin
    column_width = available_width // 2
    for column in table.columns:
        column.width = column_width
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.width = column_width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _set_cell_margins(cell, top=200 if row_index == 1 else 40, start=60, bottom=40, end=60)
    _set_table_borders_none(table)

    agency_cell, national_cell = table.rows[0].cells
    agencies = header["agency"]
    for index, line in enumerate(agencies):
        last_line = index == len(agencies) - 1
        agency_rule_width = min(3.8, max(2.0, len(line["text"]) * 0.09))
        _append_cell_line(
            agency_cell,
            line["text"],
            12,
            bold=last_line,
            bottom_rule=last_line,
            rule_width_cm=agency_rule_width if last_line else None,
        )
    _append_cell_line(national_cell, header["national"]["text"], 12, bold=True, character_scale=90)
    _append_cell_line(national_cell, header["tagline"]["text"], 13, bold=True, bottom_rule=True, rule_width_cm=5.5)

    number_cell, date_cell = table.rows[1].cells
    for line in header["number"]:
        _append_cell_line(number_cell, line["text"], 13)
    for line_index, line in enumerate(header["subject"]):
        _append_cell_line(number_cell, line["text"], 12, space_before=3 if line_index == 0 else 0)
    for line in header["date"]:
        _append_cell_line(date_cell, line["text"], 13, italic=True)

def _detect_administrative_footer(blocks: list[dict], page_width: float, page_height: float) -> dict | None:
    text_blocks = []
    for block in blocks:
        if block.get("type") != 0:
            continue
        block_text = " ".join(_line_text(line) for line in block.get("lines", [])).strip()
        if block_text:
            text_blocks.append((block, _plain_upper(block_text)))
    recipient = next((block for block, text in text_blocks if "NOI NHAN:" in text and block["bbox"][0] < page_width * 0.5), None)
    signature_roles = ("BO TRUONG", "THU TRUONG", "GIAM DOC", "CHU TICH", "QUYEN GIAM DOC")
    sign_title = next(
        (
            block
            for block, text in text_blocks
            if any(role in text for role in signature_roles)
            and "NOI NHAN" not in text
            and block["bbox"][0] >= page_width * 0.5
            and block["bbox"][1] > page_height * 0.5
        ),
        None,
    )
    if recipient is None or sign_title is None:
        return None

    name_candidates = [
        block
        for block, _text in text_blocks
        if block is not recipient
        and block is not sign_title
        and block["bbox"][0] >= page_width * 0.5
        and block["bbox"][1] > sign_title["bbox"][3]
        and block["bbox"][1] <= sign_title["bbox"][3] + 180
        and block["bbox"][1] < page_height * 0.9
    ]
    signer_name = min(name_candidates, key=lambda block: block["bbox"][1]) if name_candidates else None
    selected = [recipient, sign_title, *([signer_name] if signer_name else [])]
    boxes = [block["bbox"] for block in selected]
    return {
        "recipient": recipient,
        "sign_title": sign_title,
        "signer_name": signer_name,
        "block_ids": {id(block) for block in selected},
        "bbox": (
            min(float(box[0]) for box in boxes),
            min(float(box[1]) for box in boxes),
            max(float(box[2]) for box in boxes),
            max(float(box[3]) for box in boxes),
        ),
    }


def _append_source_block_to_cell(cell, block: dict, alignment: WD_ALIGN_PARAGRAPH) -> None:
    lines = block.get("lines", [])
    for line_index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if line_index == 0 and not cell.paragraphs[0].text else cell.add_paragraph()
        paragraph.alignment = alignment
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        line_spans = [span for span in line.get("spans", []) if str(span.get("text", "")).strip()]
        for span_index, span in enumerate(line_spans):
            text = str(span.get("text", ""))
            if span_index == 0:
                text = text.lstrip()
            if span_index == len(line_spans) - 1:
                text = text.rstrip()
            if not text:
                continue
            run = paragraph.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(max(8, min(16, float(span.get("size", 13)))))
            run.font.bold = bool(int(span.get("flags", 0)) & 16)
            run.font.italic = bool(int(span.get("flags", 0)) & 2)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if run._element.rPr is not None:
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _append_administrative_footer(document: DocxDocument, footer: dict, section) -> None:
    available_width = (
        cast(Any, section.page_width).inches
        - cast(Any, section.left_margin).inches
        - cast(Any, section.right_margin).inches
    )
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    left_width = Inches(available_width * 0.55)
    right_width = Inches(available_width * 0.45)
    _set_table_borders_none(table)
    for row in table.rows:
        for column_index, cell in enumerate(row.cells):
            cell.width = left_width if column_index == 0 else right_width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _set_cell_margins(cell, top=0, start=210 if column_index == 0 else 0, bottom=0, end=0)
    table.columns[0].width = left_width
    table.columns[1].width = right_width

    _append_source_block_to_cell(table.cell(0, 0), footer["recipient"], WD_ALIGN_PARAGRAPH.LEFT)
    signer_cell = table.cell(0, 1)
    _append_source_block_to_cell(signer_cell, footer["sign_title"], WD_ALIGN_PARAGRAPH.CENTER)
    name_block = footer.get("signer_name")
    if name_block is not None:
        gap = max(0.0, float(name_block["bbox"][1]) - float(footer["sign_title"]["bbox"][3]))
        spacer = signer_cell.add_paragraph()
        spacer.paragraph_format.space_before = Pt(gap)
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(1)
        _append_source_block_to_cell(signer_cell, name_block, WD_ALIGN_PARAGRAPH.CENTER)


def _set_page_geometry(section, width_pt: float, height_pt: float, page) -> None:
    section.page_width = Inches(max(3.0, width_pt / 72))
    section.page_height = Inches(max(3.0, height_pt / 72))
    section.orientation = WD_ORIENT.LANDSCAPE if width_pt > height_pt else WD_ORIENT.PORTRAIT
    blocks = [b for b in page.get_text("dict").get("blocks", []) if b.get("type") == 0]
    if blocks:
        left = min(max(10.0, float(b["bbox"][0])) for b in blocks)
        top = min(max(10.0, float(b["bbox"][1])) for b in blocks)
        right_edge = max(min(width_pt - 10.0, float(b["bbox"][2])) for b in blocks)
        bottom_edge = max(min(height_pt - 10.0, float(b["bbox"][3])) for b in blocks)
        section.left_margin = Inches(max(0.15, left / 72))
        section.top_margin = Inches(max(0.15, top / 72))
        section.right_margin = Inches(max(0.15, (width_pt - right_edge) / 72))
        section.bottom_margin = Inches(max(0.15, (height_pt - bottom_edge) / 72))
    else:
        section.left_margin = section.right_margin = Inches(0.55)
        section.top_margin = section.bottom_margin = Inches(0.55)


def _extract_tables(page) -> list[dict]:
    found = []
    for table in page.find_tables():
        rows = table.extract()
        if rows and any(any(cell for cell in row) for row in rows):
            cell_bboxes = [
                [tuple(float(value) for value in cell) if cell is not None else None for cell in row.cells]
                for row in table.rows
            ]
            column_count = max((len(row) for row in rows), default=0)
            x_edges = sorted(
                {
                    round(float(edge), 2)
                    for row in table.rows
                    for cell in row.cells
                    if cell
                    for edge in (cell[0], cell[2])
                }
            )
            column_widths = [right - left for left, right in zip(x_edges, x_edges[1:])]
            if len(column_widths) != column_count:
                column_widths = []
            merge_ranges: list[tuple[int, int, int]] = []
            for row_index, boxes in enumerate(cell_bboxes):
                column_index = 0
                while column_index < min(len(boxes), column_count):
                    box = boxes[column_index]
                    if box is None:
                        column_index += 1
                        continue
                    end_column = column_index
                    while end_column + 1 < min(len(boxes), column_count) and boxes[end_column + 1] is None:
                        end_column += 1
                    if end_column > column_index:
                        expected_width = sum(column_widths[column_index : end_column + 1]) if column_widths else 0
                        actual_width = float(box[2]) - float(box[0])
                        if not expected_width or abs(actual_width - expected_width) <= max(3.0, expected_width * 0.08):
                            merge_ranges.append((row_index, column_index, end_column))
                    column_index = end_column + 1
            header_row = bool(rows[0]) and any(
                token in " ".join(str(cell or "").casefold() for cell in rows[0])
                for token in (
                    "stt", "số thứ tự", "mã", "ma_", "trường hợp", "quy định", "mức hưởng",
                    "loại nhiên liệu", "tên đơn vị", "họ tên", "nội dung", "đơn vị tính", "chỉ tiêu",
                )
            )
            found.append({
                "bbox": tuple(float(x) for x in table.bbox),
                "rows": rows,
                "column_widths": column_widths,
                "cell_bboxes": cell_bboxes,
                "merge_ranges": merge_ranges,
                "header_row": header_row,
            })
    return found


def _estimate_table_font_size(blocks: list[dict], table_bbox: tuple[float, ...]) -> float | None:
    """Estimate a table's source font from text spans whose centers lie inside its grid."""
    x0, y0, x1, y1 = table_bbox
    sizes: list[float] = []
    for block in blocks:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                if not str(span.get("text", "")).strip() or not span.get("bbox"):
                    continue
                span_x0, span_y0, span_x1, span_y1 = (float(value) for value in span["bbox"])
                center_x = (span_x0 + span_x1) / 2
                center_y = (span_y0 + span_y1) / 2
                if x0 <= center_x <= x1 and y0 <= center_y <= y1:
                    sizes.append(float(span.get("size", 0)))
    if not sizes:
        return None
    sizes.sort()
    return sizes[len(sizes) // 2]


def _extract_table_cell_styles(page: Any, blocks: list[dict], cell_bboxes: list[list[tuple[float, ...] | None]]) -> list[list[dict[str, Any]]]:
    drawings = page.get_drawings()
    styles: list[list[dict[str, Any]]] = []
    for row in cell_bboxes:
        styled_row: list[dict[str, Any]] = []
        for bbox in row:
            style: dict[str, Any] = {}
            if bbox is None:
                styled_row.append(style)
                continue
            x0, y0, x1, y1 = bbox
            spans: list[dict] = []
            for block in blocks:
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        span_bbox = span.get("bbox")
                        if not span_bbox or not str(span.get("text", "")).strip():
                            continue
                        cx = (float(span_bbox[0]) + float(span_bbox[2])) / 2
                        cy = (float(span_bbox[1]) + float(span_bbox[3])) / 2
                        if x0 - 1 <= cx <= x1 + 1 and y0 - 1 <= cy <= y1 + 1:
                            spans.append(span)
            if spans:
                bold_count = sum(bool(int(span.get("flags", 0)) & 16) or "bold" in str(span.get("font", "")).casefold() for span in spans)
                italic_count = sum(bool(int(span.get("flags", 0)) & 2) or "italic" in str(span.get("font", "")).casefold() for span in spans)
                style["bold"] = bold_count > len(spans) / 2
                style["italic"] = italic_count > len(spans) / 2
                colors: dict[int, int] = {}
                for span in spans:
                    color = int(span.get("color", 0))
                    colors[color] = colors.get(color, 0) + len(str(span.get("text", "")).strip())
                if colors:
                    color = sorted(colors.items(), key=lambda pair: pair[1])[-1][0]
                    if color != 0:
                        style["color"] = f"{color & 0xFFFFFF:06X}"
                left = min(float(span["bbox"][0]) for span in spans if span.get("bbox"))
                right = max(float(span["bbox"][2]) for span in spans if span.get("bbox"))
                center = (left + right) / 2
                if right - left < (x1 - x0) * 0.75 and abs(center - (x0 + x1) / 2) <= (x1 - x0) * 0.08:
                    style["alignment"] = "center"
                elif right >= x1 - 4 and left > x0 + (x1 - x0) * 0.35:
                    style["alignment"] = "right"
                else:
                    style["alignment"] = "left"
            cell_area = max(1.0, (x1 - x0) * (y1 - y0))
            for drawing in drawings:
                fill = drawing.get("fill")
                rect = drawing.get("rect")
                if not fill or not rect or len(fill) < 3:
                    continue
                ix0, iy0 = max(x0, float(rect.x0)), max(y0, float(rect.y0))
                ix1, iy1 = min(x1, float(rect.x1)), min(y1, float(rect.y1))
                coverage = max(0.0, ix1 - ix0) * max(0.0, iy1 - iy0) / cell_area
                rgb = tuple(max(0, min(255, round(float(channel) * 255))) for channel in fill[:3])
                if coverage >= 0.85 and max(rgb) - min(rgb) > 1 and sum(rgb) < 735:
                    style["fill"] = "".join(f"{channel:02X}" for channel in rgb)
                    break
            styled_row.append(style)
        styles.append(styled_row)
    return styles


def _inside_any_table(bbox, tables: list[dict]) -> bool:
    cx = (float(bbox[0]) + float(bbox[2])) / 2
    cy = (float(bbox[1]) + float(bbox[3])) / 2
    return any(x0 - 2 <= cx <= x1 + 2 and top - 2 <= cy <= bottom + 2 for x0, top, x1, bottom in (t["bbox"] for t in tables))


def _filter_text_block_outside_tables(block: dict, tables: list[dict]) -> dict | None:
    """Keep text outside grids even when one PDF text block spans a table and a title."""
    original_lines = block.get("lines", [])
    remaining = [
        line
        for line in original_lines
        if any(str(span.get("text", "")).strip() for span in line.get("spans", []))
        and not _inside_any_table(line.get("bbox", (0, 0, 0, 0)), tables)
    ]
    if not remaining:
        return None
    if len(remaining) == len(original_lines):
        return block
    boxes = [line["bbox"] for line in remaining if line.get("bbox")]
    if not boxes:
        return None
    return {
        **block,
        "lines": remaining,
        "bbox": (
            min(float(box[0]) for box in boxes),
            min(float(box[1]) for box in boxes),
            max(float(box[2]) for box in boxes),
            max(float(box[3]) for box in boxes),
        ),
    }


def _append_table(
    document: DocxDocument,
    rows: list[list[str | None]],
    *,
    administrative: bool = False,
    column_widths: list[float] | None = None,
    available_width_pt: float | None = None,
    source_font_size: float | None = None,
    preserve_pdf_line_breaks: bool = False,
    cell_styles: list[list[dict[str, Any]]] | None = None,
    merge_ranges: list[tuple[int, int, int]] | None = None,
    existing_table: Any | None = None,
    repeat_header: bool = False,
) -> Any | None:
    width = max((len(row) for row in rows), default=0)
    if not rows or width == 0:
        return existing_table
    is_new_table = existing_table is None
    table = existing_table if existing_table is not None else document.add_table(rows=len(rows), cols=width)
    if not is_new_table:
        for _ in rows:
            table.add_row()
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    first_new_row = 0 if is_new_table else len(table.rows) - len(rows)
    if column_widths and len(column_widths) == width and sum(column_widths) > 0:
        scale = 1.0
        if available_width_pt and available_width_pt > 0:
            scale = min(1.0, available_width_pt / sum(column_widths))
        for col_index, source_width in enumerate(column_widths):
            column_width = Pt(max(1.0, source_width * scale))
            table.columns[col_index].width = column_width
            for row in table.rows:
                row.cells[col_index].width = column_width
    for row_index, row in enumerate(rows):
        for col_index in range(width):
            cell = row[col_index] if col_index < len(row) else ""
            cell_text = str(cell or "").strip()
            if not preserve_pdf_line_breaks:
                cell_text = " ".join(cell_text.split())
            word_cell = table.cell(first_new_row + row_index, col_index)
            word_cell.text = cell_text
            cell_style = cell_styles[row_index][col_index] if cell_styles and row_index < len(cell_styles) and col_index < len(cell_styles[row_index]) else {}
            if cell_style.get("fill"):
                cell_properties = word_cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), str(cell_style["fill"]))
                cell_properties.append(shading)
            alignment = cell_style.get("alignment")
            alignment_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
            for paragraph in word_cell.paragraphs:
                if alignment in alignment_map:
                    paragraph.alignment = alignment_map[alignment]
                if administrative:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    if administrative:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(max(7, min(13, source_font_size or 13)))
                        if run._element.rPr is not None:
                            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    if "bold" in cell_style:
                        run.bold = bool(cell_style["bold"])
                    if "italic" in cell_style:
                        run.italic = bool(cell_style["italic"])
                    if cell_style.get("color"):
                        run.font.color.rgb = RGBColor.from_string(str(cell_style["color"]))

    for row_index, start_column, end_column in merge_ranges or []:
        if first_new_row <= first_new_row + row_index < first_new_row + len(rows) and 0 <= start_column < end_column < width:
            table.cell(first_new_row + row_index, start_column).merge(table.cell(first_new_row + row_index, end_column))

    for row in table.rows[first_new_row : first_new_row + len(rows)]:
        row_properties = row._tr.get_or_add_trPr()
        if not row_properties.xpath("./w:cantSplit"):
            row_properties.append(OxmlElement("w:cantSplit"))
    if is_new_table and repeat_header and table.rows:
        header_properties = table.rows[0]._tr.get_or_add_trPr()
        if not header_properties.xpath("./w:tblHeader"):
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            header_properties.append(header)
    return table


def _can_continue_table(previous: dict[str, Any] | None, current: dict[str, Any], page_index: int, page_width: float, page_height: float) -> bool:
    if not previous or int(previous["page_index"]) != page_index - 1:
        return False
    previous_table = previous["source_table"]
    if max((len(row) for row in previous_table.get("rows", [])), default=0) != max((len(row) for row in current.get("rows", [])), default=0):
        return False
    previous_bbox = previous_table["bbox"]
    current_bbox = current["bbox"]
    previous_height = float(previous["page_height"])
    if float(previous_bbox[3]) < previous_height * 0.75 or float(current_bbox[1]) > page_height * 0.25:
        return False
    previous_width = float(previous["page_width"])
    if abs(float(previous_bbox[0]) / previous_width - float(current_bbox[0]) / page_width) > 0.025:
        return False
    if abs(float(previous_bbox[2]) / previous_width - float(current_bbox[2]) / page_width) > 0.025:
        return False
    old_widths = previous_table.get("column_widths") or []
    new_widths = current.get("column_widths") or []
    if old_widths and new_widths and len(old_widths) == len(new_widths):
        old_total, new_total = sum(old_widths), sum(new_widths)
        if old_total and new_total:
            deltas = [abs(old / old_total - new / new_total) for old, new in zip(old_widths, new_widths)]
            if max(deltas, default=0) > 0.08:
                return False
    return True


def _same_table_header(previous: dict[str, Any], current: dict[str, Any]) -> bool:
    if not previous.get("header_row") or not current.get("header_row"):
        return False
    old = [" ".join(str(cell or "").casefold().split()) for cell in previous["rows"][0]]
    new = [" ".join(str(cell or "").casefold().split()) for cell in current["rows"][0]]
    return old == new


def _table_has_only_side_annotations_after(events: list[tuple[float, int, dict]], table_event_index: int, table: dict) -> bool:
    x0, y0, x1, y1 = (float(value) for value in table["bbox"])
    for _top, kind, item in events[table_event_index + 1 :]:
        if kind != 1 or item.get("type") != 0:
            return False
        text = " ".join(_line_text(line) for line in item.get("lines", []) if _line_text(line)).strip()
        bbox = item.get("bbox", (0, 0, 0, 0))
        side_note = float(bbox[0]) >= x1 - 1 or float(bbox[2]) <= x0 + 1
        vertical_overlap = float(bbox[1]) <= y1 and float(bbox[3]) >= y0
        if len(text) > 40 or not side_note or not vertical_overlap:
            return False
    return True


def _append_vertical_gap(document: DocxDocument, gap_pt: float) -> None:
    if gap_pt < 2:
        return
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(gap_pt)
    run = spacer.add_run("\u00a0")
    run.font.size = Pt(1)


def _infer_text_alignment(lines: list[dict], page_width: float, block_bbox) -> WD_ALIGN_PARAGRAPH | None:
    """Infer alignment from line edges; a block's overall center is not paragraph alignment."""
    boxes = []
    for line in lines:
        box = line.get("bbox")
        if isinstance(box, (tuple, list)) and len(box) >= 4:
            boxes.append(box)
    if not page_width or not boxes:
        return None
    lefts = [float(box[0]) for box in boxes]
    rights = [float(box[2]) for box in boxes]
    centers = [(left + right) / 2 for left, right in zip(lefts, rights)]
    widths = [right - left for left, right in zip(lefts, rights)]
    page_center = page_width / 2

    if len(boxes) >= 2:
        center_error = sum(abs(center - page_center) for center in centers) / len(centers)
        left_spread = max(lefts) - min(lefts)
        right_spread = max(rights) - min(rights)
        is_centered_set = center_error <= page_width * 0.025 and max(widths) <= page_width * 0.72
        if is_centered_set and left_spread >= page_width * 0.025 and right_spread >= page_width * 0.025:
            return WD_ALIGN_PARAGRAPH.CENTER
        if (
            right_spread <= page_width * 0.025
            and left_spread >= page_width * 0.05
            and sum(lefts) / len(lefts) >= page_width * 0.40
        ):
            return WD_ALIGN_PARAGRAPH.RIGHT

    x0, _y0, x1, _y1 = block_bbox
    block_width = float(x1) - float(x0)
    block_center = (float(x0) + float(x1)) / 2
    if block_width <= page_width * 0.55 and abs(block_center - page_center) <= page_width * 0.025:
        return WD_ALIGN_PARAGRAPH.CENTER
    if float(x0) >= page_width * 0.62:
        return WD_ALIGN_PARAGRAPH.RIGHT
    return None


def _append_text_block(
    document: DocxDocument,
    block: dict,
    page_width: float,
    *,
    administrative: bool = False,
    section_left_pt: float = 0.0,
    section_right_pt: float = 0.0,
    preserve_pdf_line_breaks: bool = False,
) -> list[Any]:
    lines = block.get("lines", [])
    lines = [
        line
        for line in lines
        if any(str(span.get("text", "")).strip() for span in line.get("spans", []))
    ]
    if not lines:
        return []
    groups = split_lines_into_paragraphs(lines)
    paragraphs: list[Any] = []

    for group in groups:
        para = document.add_paragraph()
        paragraphs.append(para)
        group_boxes = [
            cast(tuple[float, float, float, float], line["bbox"])
            for line in group
            if line.get("bbox") is not None
        ]
        group_bbox = (
            min(float(box[0]) for box in group_boxes),
            min(float(box[1]) for box in group_boxes),
            max(float(box[2]) for box in group_boxes),
            max(float(box[3]) for box in group_boxes),
        ) if group_boxes else block.get("bbox", (0, 0, 0, 0))
        inferred_alignment = _infer_text_alignment(group, page_width, group_bbox)
        # Source line positions already encode paragraph gaps; do not add a second fixed gap.
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0 if preserve_pdf_line_breaks else (4 if administrative else 3))
        para.paragraph_format.line_spacing = 1.0
        if administrative:
            para.alignment = inferred_alignment or WD_ALIGN_PARAGRAPH.JUSTIFY
            min_left = float(group_bbox[0])
            first_left = float(group[0]["bbox"][0])
            para.paragraph_format.left_indent = Pt(max(0.0, min_left - section_left_pt))
            para.paragraph_format.first_line_indent = Pt(max(0.0, first_left - min_left))
            # Do not narrow the editable line width to the PDF block bounding box.
            para.paragraph_format.right_indent = Pt(0)
        elif inferred_alignment is not None:
            para.alignment = inferred_alignment
        for line_index, line in enumerate(group):
            line_spans = [s for s in line.get("spans", []) if str(s.get("text", ""))]
            for span_index, span in enumerate(line_spans):
                span_text = span.get("text", "")
                if administrative and span_index == 0:
                    span_text = span_text.lstrip()
                if administrative and span_index == len(line_spans) - 1:
                    span_text = span_text.rstrip()
                run = para.add_run(span_text)
                run.font.name = "Times New Roman" if administrative else _word_font(span.get("font", "Arial"))
                source_size = float(span.get("size", 13 if administrative else 10))
                run.font.size = Pt(max(8, min(13, source_size)) if administrative else max(5, min(48, source_size)))
                if administrative:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    if run._element.rPr is not None:
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                flags = int(span.get("flags", 0))
                run.bold = bool(flags & 16)
                run.italic = bool(flags & 2)
            if line_index < len(group) - 1:
                # Preserve deliberate heading breaks only, never source line wraps in body text.
                preserve_this_break = preserve_pdf_line_breaks and inferred_alignment == WD_ALIGN_PARAGRAPH.CENTER
                if preserve_this_break:
                    para.add_run().add_break(WD_BREAK.LINE)
                else:
                    current_text = _line_text(line).rstrip()
                    next_text = _line_text(group[line_index + 1]).lstrip()
                    if current_text and next_text and not current_text[-1].isspace() and not next_text[0].isspace():
                        para.add_run(" ")
    return paragraphs


def _word_font(font_name: str) -> str:
    name = font_name.split("+")[-1].split(",")[0].strip()
    for marker in ("-Bold", "-Italic", "Bold", "Italic", "Regular"):
        name = name.replace(marker, "")
    return name[:40] or "Arial"


def _append_image(
    document: DocxDocument,
    block: dict,
    page_width_pt: float,
    page_height_pt: float,
    z_order: int,
) -> None:
    source_bytes = block.get("image")
    if not source_bytes:
        return
    try:
        image_bytes = source_bytes
        mask_bytes = block.get("mask")
        if mask_bytes:
            image = Image.open(io.BytesIO(source_bytes)).convert("RGBA")
            mask = Image.open(io.BytesIO(mask_bytes)).convert("L")
            if mask.size != image.size:
                mask = mask.resize(image.size, Image.Resampling.LANCZOS)
            image.putalpha(ImageChops.multiply(image.getchannel("A"), mask))
            buffer = io.BytesIO()
            image.save(buffer, format="PNG")
            image_bytes = buffer.getvalue()

        bbox = block.get("bbox", (0, 0, page_width_pt, page_height_pt))
        x0, y0, x1, y1 = (float(value) for value in bbox)
        display_width = max(1.0, x1 - x0)
        display_height = max(1.0, y1 - y0)
        scale = min(1.0, page_width_pt / display_width, page_height_pt / display_height)
        display_width *= scale
        display_height *= scale
        x0 = min(max(0.0, x0), max(0.0, page_width_pt - display_width))
        y0 = min(max(0.0, y0), max(0.0, page_height_pt - display_height))
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(1)
        picture = paragraph.add_run().add_picture(
            io.BytesIO(image_bytes),
            width=Pt(display_width),
            height=Pt(display_height),
        )
        _anchor_picture(picture, x0, y0, display_width, display_height, z_order)
    except Exception:
        logger.exception("Could not preserve PDF image block")


def _anchor_picture(picture, x_pt: float, y_pt: float, width_pt: float, height_pt: float, z_order: int) -> None:
    """Turn an inline picture into a non-wrapping image at its source page coordinates."""
    inline = picture._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in (
        ("simplePos", "0"),
        ("relativeHeight", str(251658240 + z_order)),
        ("behindDoc", "0"),
        ("locked", "0"),
        ("layoutInCell", "1"),
        ("allowOverlap", "1"),
        ("distT", "0"),
        ("distB", "0"),
        ("distL", "0"),
        ("distR", "0"),
    ):
        anchor.set(name, value)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "page")
    horizontal_offset = OxmlElement("wp:posOffset")
    horizontal_offset.text = str(round(x_pt * 12700))
    position_h.append(horizontal_offset)
    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "page")
    vertical_offset = OxmlElement("wp:posOffset")
    vertical_offset.text = str(round(y_pt * 12700))
    position_v.append(vertical_offset)

    extent = inline.find(qn("wp:extent"))
    doc_properties = inline.find(qn("wp:docPr"))
    frame_properties = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    if extent is None or doc_properties is None or frame_properties is None or graphic is None:
        raise ValueError("PDF picture drawing is missing required OOXML children")
    extent.set("cx", str(round(width_pt * 12700)))
    extent.set("cy", str(round(height_pt * 12700)))
    effect_extent = OxmlElement("wp:effectExtent")
    for edge in ("l", "t", "r", "b"):
        effect_extent.set(edge, "0")
    wrap_none = OxmlElement("wp:wrapNone")
    for child in (simple_pos, position_h, position_v, extent, effect_extent, wrap_none, doc_properties, frame_properties, graphic):
        anchor.append(child)
    inline.getparent().replace(inline, anchor)


def _estimate_skew_degrees(image: Image.Image) -> float:
    """Estimate a small text-line skew and return the correction angle for PIL.rotate()."""
    probe = image.copy()
    probe.thumbnail((900, 1200))
    gray = np.asarray(probe.convert("L"))
    if gray.size == 0:
        return 0.0
    threshold = max(120, min(210, int(float(np.percentile(gray, 45)))))
    dark = gray < threshold
    height, width = dark.shape
    dark = dark[int(height * 0.08) : int(height * 0.92), int(width * 0.06) : int(width * 0.94)]
    if int(dark.sum()) < 120:
        return 0.0
    binary = Image.fromarray(np.where(dark, 0, 255).astype(np.uint8))
    best_angle = 0.0
    best_score = -1.0
    baseline_score = 0.0
    for step in range(-12, 13):
        angle = step * 0.25
        rotated = binary.rotate(angle, resample=Image.Resampling.NEAREST, expand=True, fillcolor=255)
        projection = (np.asarray(rotated) < 128).sum(axis=1).astype(np.float64)
        score = float(np.var(projection))
        if step == 0:
            baseline_score = score
        if score > best_score:
            best_score, best_angle = score, angle
    if abs(best_angle) < 0.5 or best_score < baseline_score * 1.015:
        return 0.0
    return best_angle


def _append_ocr_page(document: DocxDocument, page: Any, page_index: int, page_count: int, progress: ProgressCallback | None) -> dict[str, Any]:
    rotation_applied = 0.0
    try:
        languages = pytesseract.get_languages(config="")
        if "vie" not in languages:
            raise ConversionError("Máy chủ chưa cài dữ liệu OCR tiếng Việt (vie).")
        scale = 3.0 if max(float(page.rect.width), float(page.rect.height)) <= 1400 else 2.0
        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        image = ImageOps.autocontrast(Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("L"))
        try:
            orientation = pytesseract.image_to_osd(image, output_type=pytesseract.Output.DICT, timeout=12)
            candidate_rotation = int(orientation.get("rotate", 0)) % 360
            if candidate_rotation in {90, 180, 270}:
                image = image.rotate(-candidate_rotation, expand=True, fillcolor=255)
                rotation_applied += candidate_rotation
        except Exception:
            logger.debug("OCR orientation estimation unavailable on page %s", page_index + 1, exc_info=True)
        skew_correction = _estimate_skew_degrees(image)
        if skew_correction:
            image = image.rotate(skew_correction, expand=True, fillcolor=255)
            rotation_applied += skew_correction
        data = pytesseract.image_to_data(
            image,
            lang="vie+eng",
            config="--psm 3",
            output_type=pytesseract.Output.DICT,
            timeout=45,
        )
    except ConversionError:
        raise
    except Exception as exc:
        logger.exception("Vietnamese OCR failed on page %s", page_index + 1)
        raise ConversionError("Không thể OCR trang scan. Hãy thử PDF rõ nét hơn hoặc kiểm tra cấu hình OCR tiếng Việt.") from exc

    groups: dict[tuple[int, int], list[tuple[int, str]]] = {}
    confidences: list[float] = []
    for i, raw in enumerate(data.get("text", [])):
        word = str(raw).strip()
        try:
            confidence = float(data["conf"][i])
        except (ValueError, TypeError, KeyError):
            confidence = -1
        if not word or confidence < 0:
            continue
        confidences.append(confidence)
        key = (int(data["block_num"][i]), int(data["par_num"][i]))
        groups.setdefault(key, []).append((i, word))
    for _key, words in sorted(groups.items(), key=lambda pair: min(int(data["top"][i]) for i, _ in pair[1])):
        ordered = sorted(words, key=lambda pair: (int(data["line_num"][pair[0]]), int(data["word_num"][pair[0]])))
        by_line: dict[int, list[str]] = {}
        for index, word in ordered:
            by_line.setdefault(int(data["line_num"][index]), []).append(word)
        paragraph_text = " ".join(" ".join(line) for line in by_line.values()).strip()
        if paragraph_text:
            paragraph = document.add_paragraph(paragraph_text)
            paragraph.paragraph_format.space_after = Pt(5)
    if progress:
        progress(50 + int(25 * (page_index + 1) / max(page_count, 1)), f"OCR trang scan {page_index + 1}/{page_count}")
    return {
        "word_count": len(confidences),
        "mean_confidence": sum(confidences) / len(confidences) if confidences else None,
        "orientation_correction_degrees": round(rotation_applied, 2),
    }


def _verify_docx(path: Path) -> None:
    if not path.exists() or path.stat().st_size == 0:
        raise ConversionError("Tệp Word được tạo ra bị rỗng.")
    try:
        Document(str(path))
    except Exception as exc:
        raise ConversionError("Tệp Word đầu ra không hợp lệ.") from exc
