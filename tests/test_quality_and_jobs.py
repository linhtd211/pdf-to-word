from __future__ import annotations

import json
import time
from pathlib import Path
from typing import Any

from docx import Document

from backend import jobs
from backend.converter import _append_table, _estimate_skew_degrees, convert_pdf

FIXTURES = Path(__file__).parent / "fixtures"


def test_quality_report_covers_native_text_tables_and_editable_structure(tmp_path: Path) -> None:
    source = FIXTURES / "03-bang-bieu.pdf"
    output = tmp_path / "table-quality.docx"
    metadata = convert_pdf(source, output, "editable")
    report = metadata["quality"]

    assert report["content_comparison"]["applicable"] is True
    assert report["content_comparison"]["recall_percent"] == 100.0
    assert report["content_comparison"]["precision_percent"] == 100.0
    assert report["word_structure"]["table_count"] >= 1
    assert report["word_structure"]["textboxes"] == 0
    assert report["word_structure"]["hard_page_breaks"] == 0
    assert report["mode"] == "editable"


def test_table_writer_preserves_merge_format_and_repeating_header() -> None:
    from docx import Document as WordDocument

    document = WordDocument()
    table = _append_table(
        document,
        [["Nhóm", ""], ["Mã", "Giá trị"]],
        cell_styles=[
            [{"bold": True, "alignment": "center", "fill": "EAF1FC"}, {}],
            [{"bold": True}, {"italic": True, "color": "17365D"}],
        ],
        merge_ranges=[(0, 0, 1)],
        repeat_header=True,
    )

    assert table is not None
    assert len(table.rows) == 2
    assert table._tbl.xpath(".//w:gridSpan")
    assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert len(table._tbl.xpath(".//w:cantSplit")) == 2
    assert table.cell(0, 0)._tc.xpath("./w:tcPr/w:shd[@w:fill='EAF1FC']")
    assert table.cell(1, 1).paragraphs[0].runs[0].italic is True
    assert str(table.cell(1, 1).paragraphs[0].runs[0].font.color.rgb) == "17365D"


def test_adjacent_pdf_page_tables_join_as_one_editable_word_table(tmp_path: Path) -> None:
    import fitz

    font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    assert font_path.exists()
    source = tmp_path / "continued-table.pdf"
    pdf: Any = fitz.open()
    rows_by_page = [
        [["STT", "Nội dung"], ["1", "Dòng đầu"], ["2", "Dòng cuối trang"]],
        [["STT", "Nội dung"], ["3", "Dòng tiếp trang"]],
    ]
    x = [55, 150, 540]
    page_y = [[650, 690, 730, 775], [35, 75, 115]]
    for page_index, page_rows in enumerate(rows_by_page):
        page = pdf.new_page(width=595.28, height=841.89)
        ys = page_y[page_index]
        for x_value in x:
            page.draw_line((x_value, ys[0]), (x_value, ys[-1]), color=(0, 0, 0), width=0.7)
        for y_value in ys:
            page.draw_line((x[0], y_value), (x[-1], y_value), color=(0, 0, 0), width=0.7)
        for row_index, row in enumerate(page_rows):
            for column_index, value in enumerate(row):
                page.insert_text(
                    (x[column_index] + 5, ys[row_index] + 22),
                    value,
                    fontname="vn",
                    fontfile=str(font_path),
                    fontsize=10,
                )
    pdf.save(source)
    pdf.close()

    output = tmp_path / "continued-table.docx"
    convert_pdf(source, output, "editable")
    document = Document(str(output))
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 4  # repeated second-page header is not duplicated
    assert document.tables[0].rows[3].cells[1].text == "Dòng tiếp trang"
    assert document.tables[0].rows[0]._tr.xpath("./w:trPr/w:tblHeader")

    administrative_output = tmp_path / "continued-table-admin.docx"
    convert_pdf(source, administrative_output, "layout", page_policy="administrative")
    administrative_document = Document(str(administrative_output))
    assert len(administrative_document.tables) == 1
    assert len(administrative_document.tables[0].rows) == 4
    assert not administrative_document._element.xpath(".//w:br[@w:type='page']")


def test_recovery_marks_missing_inflight_upload_as_interrupted(tmp_path: Path, monkeypatch) -> None:
    job_root = tmp_path / "jobs"
    workdir = job_root / "interrupted"
    workdir.mkdir(parents=True)
    job_id = "interrupted-after-restart"
    manifest = {
        "id": job_id,
        "status": "processing",
        "progress": 45,
        "stage": "Đang xử lý",
        "created_at": time.time(),
        "updated_at": time.time(),
        "workdir": str(workdir),
        "source": str(workdir / "source.pdf"),
        "output": str(workdir / "result.docx"),
        "download_name": "sample.docx",
        "mode": "editable",
        "page_policy": "source",
        "pages": 1,
        "size": 0,
        "warnings": [],
        "error": None,
    }
    (workdir / "job.json").write_text(json.dumps(manifest), encoding="utf-8")
    monkeypatch.setattr(jobs, "JOB_ROOT", job_root)

    assert jobs.recover_jobs() == 1
    restored = jobs.get_job(job_id)
    assert restored is not None
    assert restored["status"] == "failed"
    assert "khởi động lại" in restored["error"].lower()
    jobs.cleanup_job(job_id)


def test_ocr_deskews_low_resolution_vietnamese_scan_and_reports_confidence(tmp_path: Path) -> None:
    import io

    import fitz
    from PIL import Image, ImageDraw, ImageFont

    font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    font = ImageFont.truetype(str(font_path), 26)
    scan = Image.new("L", (900, 1200), 255)
    draw = ImageDraw.Draw(scan)
    for line in range(16):
        draw.text((65, 45 + line * 66), "Quyết định và văn bản hành chính tiếng Việt 0123456789", font=font, fill=0)
    tilted = scan.rotate(2.0, resample=Image.Resampling.BICUBIC, expand=False, fillcolor=255)
    correction = _estimate_skew_degrees(tilted)
    assert abs(correction + 2.0) <= 0.5

    image_buffer = io.BytesIO()
    tilted.save(image_buffer, format="PNG")
    source = tmp_path / "tilted-scan.pdf"
    pdf: Any = fitz.open()
    page = pdf.new_page(width=595.28, height=841.89)
    page.insert_image(page.rect, stream=image_buffer.getvalue())
    pdf.save(source)
    pdf.close()

    output = tmp_path / "tilted-scan.docx"
    metadata = convert_pdf(source, output, "editable")
    ocr = metadata["quality"]["ocr"]
    assert len(ocr["pages"]) == 1
    assert abs(float(ocr["pages"][0]["orientation_correction_degrees"]) + 2.0) <= 0.5
    assert ocr["pages"][0]["word_count"] > 20
    assert "scan" in " ".join(metadata["warnings"]).casefold()


def test_conversion_queue_has_a_hard_capacity_and_releases_reserved_slots() -> None:
    reservations = [jobs.reserve_capacity() for _ in range(jobs.MAX_ACCEPTED_JOBS)]
    try:
        assert all(reservation is not None for reservation in reservations)
        assert jobs.reserve_capacity() is None
    finally:
        for reservation in reservations:
            jobs.release_capacity(reservation)
    reservation = jobs.reserve_capacity()
    assert reservation is not None
    jobs.release_capacity(reservation)
