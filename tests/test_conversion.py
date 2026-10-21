from __future__ import annotations

import time
from pathlib import Path
from typing import Any, cast

from docx import Document
from fastapi.testclient import TestClient
import pytest

from backend.app import app
from backend.converter import ConversionError, convert_pdf, inspect_pdf

FIXTURES = Path(__file__).parent / "fixtures"


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)


def test_electronic_pdf_to_editable_docx(tmp_path):
    source = FIXTURES / "01-van-ban-dien-tu.pdf"
    info = inspect_pdf(source)
    assert info["kind"] == "electronic"
    result = tmp_path / "text.docx"
    convert_pdf(source, result, "editable")
    text = _docx_text(result)
    assert "HƯỚNG DẪN SỬ DỤNG HỆ THỐNG" in text
    assert "Độc lập" in text
    assert Document(result)


def test_table_pdf_creates_real_editable_word_table(tmp_path):
    source = FIXTURES / "03-bang-bieu.pdf"
    result = tmp_path / "table.docx"
    convert_pdf(source, result, "editable")
    doc = Document(result)
    assert len(doc.tables) >= 1
    values = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert "Tên đơn vị" in values
    assert "Phòng Hành chính" in values


def test_editable_table_cells_reflow_while_layout_cells_keep_source_breaks():
    from docx import Document as WordDocument

    from backend.converter import _append_table

    rows: list[list[str | None]] = [["Dòng đầu\nDòng sau", "100%"]]
    editable = WordDocument()
    _append_table(editable, rows, preserve_pdf_line_breaks=False)
    editable_cell = editable.tables[0].cell(0, 0)
    assert editable_cell.text == "Dòng đầu Dòng sau"
    assert not editable_cell._tc.xpath(".//w:br")

    layout = WordDocument()
    _append_table(layout, rows, preserve_pdf_line_breaks=True)
    layout_cell = layout.tables[0].cell(0, 0)
    assert layout_cell.text == "Dòng đầu\nDòng sau"
    assert len(layout_cell._tc.xpath(".//w:br")) == 1


def test_four_page_pdf_reuses_identical_section_and_flows_naturally(tmp_path):
    source = FIXTURES / "04-nhieu-trang.pdf"
    assert inspect_pdf(source)["pages"] == 4
    result = tmp_path / "multipage.docx"
    convert_pdf(source, result, "editable")
    doc = Document(result)
    assert len(doc.sections) == 1
    assert not doc._element.xpath(".//w:br[@w:type='page']")


def test_embedded_pdf_illustration_is_a_real_word_image(tmp_path):
    result = tmp_path / "image.docx"
    convert_pdf(FIXTURES / "05-co-hinh-anh.pdf", result, "editable")
    doc = Document(result)
    assert len(doc._element.xpath(".//wp:anchor")) >= 1


def test_scan_pdf_uses_vietnamese_ocr_and_produces_text(tmp_path):
    source = FIXTURES / "06-scan-tieng-viet.pdf"
    assert inspect_pdf(source)["kind"] == "scanned"
    result = tmp_path / "scan.docx"
    metadata = convert_pdf(source, result, "editable")
    text = _docx_text(result)
    assert metadata["kind"] == "scanned"
    assert len(text.strip()) > 100
    assert "THÔNG BÁO" in text or "THONG BAO" in text
    assert "scan" in metadata["warnings"][0].lower()


def test_layout_mode_explicitly_warns_on_scan_fallback(tmp_path):
    result = tmp_path / "scan-layout.docx"
    metadata = convert_pdf(FIXTURES / "06-scan-tieng-viet.pdf", result, "layout")
    assert any("chưa thể tái tạo đầy đủ" in warning for warning in metadata["warnings"])


def test_layout_mode_pdf2docx_returns_openable_docx(tmp_path):
    result = tmp_path / "layout.docx"
    convert_pdf(FIXTURES / "01-van-ban-dien-tu.pdf", result, "layout")
    assert result.stat().st_size > 1000
    assert len(Document(result).paragraphs) > 0


def test_administrative_header_uses_balanced_two_column_table_in_both_modes(tmp_path):
    source = FIXTURES / "02-cong-van-hanh-chinh.pdf"
    for mode in ("editable", "layout"):
        result = tmp_path / f"admin-{mode}.docx"
        metadata = convert_pdf(source, result, mode)
        doc = Document(str(result))
        assert metadata["warnings"]
        assert len(doc.tables) == 1
        header = doc.tables[0]
        assert len(header.rows) == 2
        assert len(header.columns) == 2
        assert header.columns[0].width == header.columns[1].width
        left_top = header.cell(0, 0).text
        right_top = header.cell(0, 1).text
        left_bottom = header.cell(1, 0).text
        right_bottom = header.cell(1, 1).text
        assert "SỞ NỘI VỤ" in left_top
        assert "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in right_top
        assert "Độc lập - Tự do - Hạnh phúc" in right_top
        tagline_paragraph = header.cell(0, 1).paragraphs[1]
        assert tagline_paragraph.text == "Độc lập - Tự do - Hạnh phúc"
        assert tagline_paragraph.paragraph_format.left_indent is None
        assert tagline_paragraph.paragraph_format.right_indent is None
        assert header._tbl.xml.count("<w:pBdr>") == 2
        assert "Số: 15/CV-SNV" in left_bottom
        assert "V/v hướng dẫn lập hồ sơ" in left_bottom
        assert "Hà Nội, ngày 15 tháng 10 năm 2026" in right_bottom
        sections = cast(Any, doc.sections)
        assert all(section.page_width.inches == pytest.approx(8.2677, abs=0.01) for section in sections)
        assert all(section.left_margin.cm == pytest.approx(3.0, abs=0.01) for section in sections)
        assert all(section.right_margin.cm == pytest.approx(1.5, abs=0.01) for section in sections)
        normal_style = cast(Any, doc.styles["Normal"])
        assert normal_style.font.name == "Times New Roman"
        assert not any("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in para.text for para in doc.paragraphs)


def test_administrative_custom_landscape_preserves_page_size_and_table_font(tmp_path):
    import fitz
    from docx.enum.section import WD_ORIENT

    font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    assert font_path.exists()
    source = tmp_path / "admin-custom-landscape.pdf"
    pdf: Any = fitz.open()
    portrait = pdf.new_page(width=595.28, height=841.89)
    text_options = {"fontname": "vn", "fontfile": str(font_path), "fontsize": 12}
    portrait.insert_text((55, 48), "BỘ Y TẾ", **text_options)
    portrait.insert_text((310, 48), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", **text_options)
    portrait.insert_text((330, 66), "Độc lập - Tự do - Hạnh phúc", **text_options)
    portrait.insert_text((55, 120), "Nội dung quyết định.", **text_options)

    landscape = pdf.new_page(width=905.5, height=612)
    widths = [45, 120, 210, 190, 150]
    row_heights = [28, 36]
    x_positions = [20]
    for width in widths:
        x_positions.append(x_positions[-1] + width)
    y_positions = [80]
    for height in row_heights:
        y_positions.append(y_positions[-1] + height)
    for x in x_positions:
        landscape.draw_line((x, y_positions[0]), (x, y_positions[-1]), color=(0, 0, 0), width=0.5)
    for y in y_positions:
        landscape.draw_line((x_positions[0], y), (x_positions[-1], y), color=(0, 0, 0), width=0.5)
    values = [["STT", "Mã", "Trường hợp", "Quy định", "Mức hưởng"], ["1", "1.1", "Nội dung bảng", "Căn cứ mẫu", "100%"]]
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            landscape.insert_text(
                (x_positions[column_index] + 3, y_positions[row_index] + 18),
                value,
                **text_options,
            )
    pdf.save(source)
    pdf.close()

    output = tmp_path / "admin-custom-landscape.docx"
    convert_pdf(source, output, "editable")
    doc = Document(str(output))
    assert len(doc.sections) == 2
    landscape_section = doc.sections[1]
    page_width = landscape_section.page_width
    page_height = landscape_section.page_height
    assert page_width is not None and page_width.pt == pytest.approx(905.5, abs=0.2)
    assert page_height is not None and page_height.pt == pytest.approx(612, abs=0.2)
    assert landscape_section.orientation == WD_ORIENT.LANDSCAPE
    appendix = next(table for table in doc.tables if len(table.columns) == 5)
    assert len(appendix.rows) == 2
    table_runs = [run for row in appendix.rows for cell in row.cells for paragraph in cell.paragraphs for run in paragraph.runs]
    assert table_runs and all(run.font.size is not None and run.font.size.pt == pytest.approx(12) for run in table_runs)


def test_administrative_group_indentation_uses_the_full_source_block_width():
    from docx import Document as WordDocument
    from docx.shared import Pt

    from backend.converter import _append_text_block

    document = WordDocument()
    block = {
        "bbox": (80, 100, 310, 150),
        "lines": [
            {
                "bbox": (80, 100, 135, 115),
                "spans": [
                    {"text": "Nơi", "size": 12, "font": "Times New Roman"},
                    {"text": " ", "size": 12, "font": "Times New Roman"},
                    {"text": "nhận:", "size": 12, "font": "Times New Roman"},
                ],
            },
            {"bbox": (80, 122, 170, 137), "spans": [{"text": "- Mục nhận:", "size": 12, "font": "Times New Roman"}]},
        ],
    }
    _append_text_block(
        document,
        block,
        page_width=595,
        administrative=True,
        section_left_pt=85,
        section_right_pt=42.5,
        preserve_pdf_line_breaks=True,
    )
    expected_right_indent = Pt(595 - 42.5 - 310)
    assert len(document.paragraphs) == 2
    assert document.paragraphs[0].text == "Nơi nhận:"
    assert all(paragraph.paragraph_format.right_indent == expected_right_indent for paragraph in document.paragraphs)
    assert all(paragraph.paragraph_format.space_before in (None, Pt(0)) for paragraph in document.paragraphs)


def test_editable_mode_reflows_pdf_lines_as_a_natural_word_paragraph():
    from docx import Document as WordDocument

    from backend.converter import _append_text_block

    block = {
        "bbox": (72, 100, 300, 126),
        "lines": [
            {"bbox": (72, 100, 300, 112), "spans": [{"text": "Nội dung đoạn văn tiếp tục trên"}]},
            {"bbox": (72, 114, 210, 126), "spans": [{"text": "dòng sau trong cùng đoạn."}]},
        ],
    }
    document = WordDocument()
    _append_text_block(document, block, page_width=595, preserve_pdf_line_breaks=False)

    assert len(document.paragraphs) == 1
    assert document.paragraphs[0].text == "Nội dung đoạn văn tiếp tục trên dòng sau trong cùng đoạn."
    assert not document.paragraphs[0]._p.xpath(".//w:br")


def test_layout_mode_keeps_pdf_line_breaks():
    from docx import Document as WordDocument

    from backend.converter import _append_text_block

    block = {
        "bbox": (72, 100, 300, 126),
        "lines": [
            {"bbox": (72, 100, 300, 112), "spans": [{"text": "Dòng đầu của mẫu PDF"}]},
            {"bbox": (72, 114, 210, 126), "spans": [{"text": "Dòng tiếp theo."}]},
        ],
    }
    document = WordDocument()
    _append_text_block(document, block, page_width=595, preserve_pdf_line_breaks=True)

    assert len(document.paragraphs) == 1
    assert len(document.paragraphs[0]._p.xpath(".//w:br")) == 1


def test_editable_administrative_paragraph_uses_full_word_width():
    from docx import Document as WordDocument

    from backend.converter import _append_text_block

    block = {
        "bbox": (80, 100, 310, 126),
        "lines": [
            {"bbox": (80, 100, 230, 112), "spans": [{"text": "Nội dung hành chính tiếp tục"}]},
            {"bbox": (80, 114, 180, 126), "spans": [{"text": "trên một dòng khác."}]},
        ],
    }
    document = WordDocument()
    _append_text_block(
        document,
        block,
        page_width=595,
        administrative=True,
        section_left_pt=85,
        section_right_pt=42.5,
        preserve_pdf_line_breaks=False,
    )

    paragraph = document.paragraphs[0]
    assert paragraph.text == "Nội dung hành chính tiếp tục trên một dòng khác."
    assert paragraph.paragraph_format.right_indent is not None
    assert paragraph.paragraph_format.right_indent.pt == 0
    assert not paragraph._p.xpath(".//w:br")


def test_invalid_pdf_is_rejected(tmp_path):
    invalid = tmp_path / "invalid.pdf"
    invalid.write_bytes(b"not a PDF")
    try:
        inspect_pdf(invalid)
    except ConversionError:
        pass
    else:
        raise AssertionError("Tệp không phải PDF phải bị từ chối")


def test_api_upload_progress_and_download():
    source = FIXTURES / "01-van-ban-dien-tu.pdf"
    with TestClient(app) as client:
        assert client.get("/api/health").status_code == 200
        response = client.post("/api/conversions", data={"mode": "editable"}, files={"file": ("Văn bản mẫu.pdf", source.read_bytes(), "application/pdf")})
        assert response.status_code == 202, response.text
        job_id = response.json()["id"]
        deadline = time.time() + 30
        payload: dict = {}
        while time.time() < deadline:
            status = client.get(f"/api/conversions/{job_id}")
            assert status.status_code == 200
            payload = status.json()
            if payload["status"] in {"completed", "failed"}:
                break
            time.sleep(.2)
        assert payload["status"] == "completed", payload
        assert payload["progress"] == 100
        assert payload["quality"]["content_comparison"]["recall_percent"] == 100.0
        assert payload["quality"]["word_structure"]["textboxes"] == 0
        download = client.get(f"/api/conversions/{job_id}/download")
        assert download.status_code == 200
        assert download.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
        assert payload["download_name"] == "Văn bản mẫu.docx"
        assert "filename*=utf-8''" in download.headers["content-disposition"].lower()
        assert len(download.content) > 1000
        assert client.get(f"/api/conversions/{job_id}").status_code == 404


def test_admin_form_preserves_transparency_landscape_and_body_alignment(tmp_path):
    import io
    from zipfile import ZipFile

    import fitz
    from PIL import Image, ImageDraw
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    assert font_path.exists()
    source = tmp_path / "mixed-admin-form.pdf"
    pdf: Any = fitz.open()
    page = pdf.new_page(width=595, height=842)
    text_options = {"fontname": "vn", "fontfile": str(font_path), "fontsize": 10}
    page.insert_text((55, 42), "SỞ Y TẾ", **text_options)
    page.insert_text((315, 42), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", **text_options)
    page.insert_text((335, 62), "Độc lập - Tự do - Hạnh phúc", **text_options)
    page.insert_text((55, 88), "Số:           /SYT", **text_options)
    page.insert_text((113, 88), "01", **text_options)
    page.insert_text((315, 88), "Hà Nội, ngày       tháng       năm 2026", **text_options)
    page.insert_text((422, 88), "01", **text_options)
    page.insert_text((486, 88), "10", **text_options)
    page.insert_text((55, 108), "V/v báo cáo kết quả hỗ trợ chuyên môn", **text_options)
    page.insert_text((55, 122), "giữa đơn vị với y tế cơ sở", **text_options)
    body = "Nội dung thân thư có nhiều dòng và cần giữ canh trái.\nDòng tiếp theo cùng lề trái, không phải tiêu đề căn giữa."
    page.insert_textbox(fitz.Rect(60, 150, 535, 260), body, **{**text_options, "fontsize": 11})

    stamp = Image.new("RGBA", (240, 240), (0, 0, 0, 0))
    stamp_draw = ImageDraw.Draw(stamp)
    stamp_draw.ellipse((10, 10, 230, 230), outline=(195, 35, 35, 255), width=8)
    stamp_draw.ellipse((28, 28, 212, 212), outline=(195, 35, 35, 255), width=3)
    stamp_buffer = io.BytesIO()
    stamp.save(stamp_buffer, format="PNG")
    page.insert_image(fitz.Rect(320, 545, 430, 655), stream=stamp_buffer.getvalue())

    signature = Image.new("RGBA", (300, 100), (0, 0, 0, 0))
    draw = ImageDraw.Draw(signature)
    draw.line((10, 75, 75, 28, 130, 68, 280, 18), fill=(20, 40, 210, 255), width=6)
    buffer = io.BytesIO()
    signature.save(buffer, format="PNG")
    page.insert_image(fitz.Rect(350, 620, 500, 670), stream=buffer.getvalue())

    page.insert_textbox(fitz.Rect(55, 575, 240, 660), "Nơi nhận:\\n- Như trên;\\n- Lưu: VT, NVY.", **text_options)
    page.insert_textbox(
        fitz.Rect(360, 575, 535, 620),
        "KT. GIÁM ĐỐC\\nPHÓ GIÁM ĐỐC",
        **{**text_options, "fontname": "vn-bold"},
    )
    page.insert_text((397, 650), "Nguyễn Văn A", **text_options)
    # An additional small transparent pen mark overlaps the stamp, as in scanned signatures.
    mark = Image.new("RGBA", (120, 50), (0, 0, 0, 0))
    ImageDraw.Draw(mark).line((5, 35, 30, 10, 65, 28, 115, 5), fill=(20, 40, 210, 230), width=4)
    mark_buffer = io.BytesIO()
    mark.save(mark_buffer, format="PNG")
    page.insert_image(fitz.Rect(385, 575, 445, 600), stream=mark_buffer.getvalue())

    landscape = pdf.new_page(width=842, height=595)

    def draw_table(x: float, y: float, widths: list[float], row_heights: list[float], rows: list[list[str]]) -> None:
        xs = [x]
        for cell_width in widths:
            xs.append(xs[-1] + cell_width)
        ys = [y]
        for row_height in row_heights:
            ys.append(ys[-1] + row_height)
        for x_pos in xs:
            landscape.draw_line((x_pos, ys[0]), (x_pos, ys[-1]), color=(0, 0, 0), width=0.5)
        for y_pos in ys:
            landscape.draw_line((xs[0], y_pos), (xs[-1], y_pos), color=(0, 0, 0), width=0.5)
        for row_index, row in enumerate(rows):
            for col_index, text in enumerate(row):
                if not text:
                    continue
                cell = fitz.Rect(xs[col_index] + 2, ys[row_index] + 2, xs[col_index + 1] - 2, ys[row_index + 1] - 2)
                landscape.insert_textbox(cell, text, fontname="vn-bold" if row_index == 0 else "vn", fontfile=str(font_path), fontsize=7, align=1 if row_index == 0 else 0)

    first_widths = [38, 112, 100, 92, 92, 78, 50, 50, 58, 50]
    first_rows = [
        ["STT", "Tên Bệnh viện", "Số TYT đã\\nhỗ trợ", "Số DMKT đã\\nchuyển giao", "Đào tạo\\nchuyên môn", "Đánh giá\\nhiệu quả", "Khó\\nkhăn", "Đề\\nxuất", "Phương\\nhướng", "Ghi\\nchú"],
        ["1", "BV ..", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", ""],
    ]
    second_widths = [37, 86, 101, 85, 85, 64, 71, 48, 47, 58, 47]
    second_rows = [
        ["STT", "Tên TYT", "Số Bệnh viện\\nđã hỗ trợ\\nchuyên môn", "Số DMKT\\nđã được\\nchuyển giao", "Đào tạo\\nnâng cao\\nchuyên môn", "Đủ điều kiện\\ncông bố KSK,\\ntham gia KSK\\ntoàn dân", "Đánh giá\\nhiệu quả", "Khó\\nkhăn", "Đề\\nxuất", "Phương\\nhướng", "Ghi\\nchú"],
        ["1", "TYT..", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", "", ""],
    ]
    first_widths = [value * 720 / sum(first_widths) for value in first_widths]
    second_widths = [value * 729 / sum(second_widths) for value in second_widths]
    landscape.insert_text((285, 86), "PHỤ LỤC I: (MẪU) BÁO CÁO BỆNH VIỆN", **{**text_options, "fontname": "vn-bold"})
    draw_table(56, 100, first_widths, [32, 18, 17, 17], first_rows)
    landscape.insert_text((284, 282), "PHỤ LỤC II: (MẪU) BÁO CÁO TRẠM Y TẾ", **{**text_options, "fontname": "vn-bold"})
    draw_table(56, 297, second_widths, [113, 17, 17, 17], second_rows)
    pdf.save(source)
    pdf.close()

    output = tmp_path / "mixed-admin-form.docx"
    convert_pdf(source, output, "editable")
    doc = Document(str(output))
    assert len(doc.sections) == 2
    national_run = doc.tables[0].rows[0].cells[1].paragraphs[0].runs[0]
    assert national_run.font.size is not None and national_run.font.size.pt == 12
    national_rpr = national_run._element.rPr
    assert national_rpr is not None
    national_scale = national_rpr.find(qn("w:w"))
    assert national_scale is not None and national_scale.get(qn("w:val")) == "90"
    landscape_section = doc.sections[1]
    assert landscape_section.orientation == WD_ORIENT.LANDSCAPE
    page_width = landscape_section.page_width
    page_height = landscape_section.page_height
    assert page_width is not None and page_height is not None
    assert page_width.inches > page_height.inches
    body_paragraph = next(p for p in doc.paragraphs if "Nội dung thân thư" in p.text)
    assert body_paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER
    anchors = doc._element.xpath(".//wp:anchor")
    assert len(anchors) == 3
    first_extent = anchors[0].find(qn("wp:extent"))
    assert first_extent is not None
    assert 100 <= int(first_extent.get("cx")) / 12700 <= 120
    first_position = anchors[0].find(qn("wp:positionV"))
    assert first_position is not None
    first_offset = first_position.find(qn("wp:posOffset"))
    assert first_offset is not None and int(first_offset.text or "0") / 12700 == pytest.approx(545)

    text = _docx_text(output)
    assert "Số: 01 /SYT" in text
    assert "ngày 01 tháng 10 năm 2026" in text
    assert "giữa đơn vị với y tế cơ sở" in text
    assert "Nơi nhận:" in text and "KT. GIÁM ĐỐC" in text and "Nguyễn Văn A" in text
    appendix_tables = [table for table in doc.tables if len(table.columns) in {10, 11}]
    assert sorted(len(table.columns) for table in appendix_tables) == [10, 11]
    assert all(len(table.rows) == 4 for table in appendix_tables)
    for table in appendix_tables:
        widths_pt = [cast(Any, column.width).pt for column in table.columns]
        assert max(widths_pt) / min(widths_pt) > 1.8
    assert sum(len(table.columns) for table in doc.tables) >= 25

    with ZipFile(output) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert len(media) == 3
        for media_name in media:
            converted = Image.open(io.BytesIO(archive.read(media_name))).convert("RGBA")
            alpha_min, alpha_max = cast(tuple[int, int], converted.getchannel("A").getextrema())
            assert alpha_min == 0 and alpha_max >= 230
