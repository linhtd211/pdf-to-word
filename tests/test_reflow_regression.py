from docx import Document
from backend.converter import _append_text_block, _append_table


def _line(text, y):
    return {"bbox": (75, y, 520, y + 14), "spans": [{"text": text, "font": "TimesNewRoman", "size": 13, "flags": 0}]}


def test_layout_paragraph_has_no_hard_line_breaks():
    doc = Document()
    block = {"bbox": (75, 80, 520, 124), "lines": [
        _line("Điều 1. Ban hành danh mục mã đối tượng đến khám bệnh, chữa bệnh và", 80),
        _line("Danh mục mã nhiên liệu phục vụ việc gửi dữ liệu điện tử chi phí", 95),
        _line("khám bệnh, chữa bệnh bảo hiểm y tế.", 110),
    ]}
    paragraphs = _append_text_block(doc, block, 595, administrative=True, section_left_pt=75, section_right_pt=40, preserve_pdf_line_breaks=False)
    assert len(paragraphs) == 1
    assert not paragraphs[0]._p.xpath('.//w:br')
    assert 'và Danh mục' in paragraphs[0].text


def test_layout_table_cell_reflows():
    doc = Document()
    _append_table(doc, [["Đây là dòng đầu\nđây là dòng tiếp theo"]], preserve_pdf_line_breaks=False)
    assert not doc.tables[0].cell(0, 0)._tc.xpath('.//w:br')
